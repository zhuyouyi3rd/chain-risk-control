from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PATH = Path("docs/新能源汽车产业链_企业关系示例.xlsx")

FONT = "Microsoft YaHei"
TEXT = "263238"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT, size=16, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10, color=TEXT)
LABEL_FONT = Font(name=FONT, size=10, bold=True, color="1F4E79")
BAND_FILL = PatternFill("solid", fgColor="F4F8FB")
LABEL_FILL = PatternFill("solid", fgColor="EAF2F8")
TRUE_FILL = PatternFill("solid", fgColor="E7F5EC")
FALSE_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="C7D6E2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

WIDTHS = {
    "填写说明": [22, 105],
    "产业链说明": [16, 10, 22, 16, 55, 50],
    "关系速览": [10, 22, 22, 14, 10, 50],
    "企业表": [10, 24, 16, 16, 48, 14, 14],
    "关系表": [10, 12, 12, 14, 12, 32, 10],
    "风险源表": [10, 26, 16, 28, 10, 10],
}

CENTER_COLUMNS = {
    "产业链说明": {2, 4},
    "关系速览": {1, 4, 5},
    "企业表": {1, 3, 6, 7},
    "关系表": {1, 2, 3, 4, 5, 7},
    "风险源表": {1, 3, 5, 6},
}

ROW_HEIGHTS = {
    "填写说明": 27,
    "产业链说明": 34,
    "关系速览": 28,
    "企业表": 30,
    "关系表": 24,
    "风险源表": 26,
}


def style_sheet(sheet):
    title = sheet.title
    show_grid = False
    sheet.sheet_view.showGridLines = show_grid
    sheet.sheet_view.zoomScale = 100
    sheet.freeze_panes = "A2"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    widths = WIDTHS.get(title, [])
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row in sheet.iter_rows():
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER

    sheet.row_dimensions[1].height = 30
    for row_index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = ROW_HEIGHTS.get(title, 26)

    if title == "填写说明":
        sheet.row_dimensions[1].height = 38
        sheet["A1"].font = TITLE_FONT
        sheet["A1"].fill = HEADER_FILL
        sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
        for row_index in range(2, sheet.max_row + 1):
            label = sheet.cell(row=row_index, column=1)
            value = sheet.cell(row=row_index, column=2)
            label.font = LABEL_FONT
            label.fill = LABEL_FILL
            label.alignment = WRAP_LEFT
            value.font = BODY_FONT
            value.fill = PatternFill("solid", fgColor="FFFFFF")
            value.alignment = WRAP_LEFT
            sheet.row_dimensions[row_index].height = 28
    else:
        for cell in sheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = CENTER
            cell.border = BORDER
        center_columns = CENTER_COLUMNS.get(title, set())
        for row_index in range(2, sheet.max_row + 1):
            for column_index in range(1, sheet.max_column + 1):
                cell = sheet.cell(row=row_index, column=column_index)
                cell.fill = BAND_FILL if row_index % 2 == 0 else PatternFill(
                    "solid", fgColor="FFFFFF"
                )
                cell.alignment = (
                    CENTER if column_index in center_columns else WRAP_LEFT
                )
        if sheet.auto_filter.ref:
            sheet.auto_filter.ref = sheet.auto_filter.ref

    if title == "关系表":
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row_index, column=5)
            cell.number_format = "0.00"
            if sheet.cell(row=row_index, column=7).value is True:
                sheet.cell(row=row_index, column=7).fill = TRUE_FILL
    elif title == "风险源表":
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row_index, column=5)
            cell.number_format = "0.00"
            enabled = sheet.cell(row=row_index, column=6)
            enabled.fill = TRUE_FILL if enabled.value else FALSE_FILL
    elif title == "关系速览":
        for row_index in range(2, sheet.max_row + 1):
            sheet.cell(row=row_index, column=5).number_format = "0.00"

    sheet.sheet_properties.tabColor = {
        "填写说明": "1F4E79",
        "产业链说明": "2E75B6",
        "关系速览": "2E75B6",
        "企业表": "70AD47",
        "关系表": "70AD47",
        "风险源表": "C00000",
    }.get(title, "808080")


workbook = load_workbook(PATH, data_only=False)
for worksheet in workbook.worksheets:
    style_sheet(worksheet)
workbook.active = 0
workbook.save(PATH)

check = load_workbook(PATH, data_only=False)
print("saved:", PATH)
for worksheet in check.worksheets:
    print(
        worksheet.title,
        worksheet.max_row,
        worksheet.max_column,
        "merged=", len(worksheet.merged_cells.ranges),
        "filter=", worksheet.auto_filter.ref,
        "freeze=", worksheet.freeze_panes,
    )

assert check["企业表"].max_row == 13
assert check["关系表"].max_row == 17
assert check["风险源表"].max_row == 4
assert check["企业表"]["A2"].value == "A01"
assert check["关系表"]["E2"].value == 0.85
assert check["风险源表"]["F2"].value is True
assert check["风险源表"]["F4"].value is False
assert check["企业表"].freeze_panes == "A2"
assert check["填写说明"]["A1"].value == "新能源汽车产业链 · 风控示例数据"
print("✓ styles and values preserved")
